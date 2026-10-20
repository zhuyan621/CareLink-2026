"""Build the new-colleague launch guide DOCX from its Markdown source."""

from __future__ import annotations

import re
from pathlib import Path

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_BREAK, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "START_HERE_新同事打开网页.md"
OUTPUT = ROOT / "新同事_从零打开网页.docx"

INK = RGBColor(28, 39, 51)
MUTED = RGBColor(87, 102, 117)
BLUE = RGBColor(21, 82, 112)
PALE = "EDF3F6"


def set_font(style, *, size: float, bold: bool = False, color=INK) -> None:
    style.font.name = "Microsoft YaHei"
    style.font.size = Pt(size)
    style.font.bold = bold
    style.font.color.rgb = color
    rpr = style.element.get_or_add_rPr()
    rfonts = rpr.rFonts
    if rfonts is None:
        rfonts = OxmlElement("w:rFonts")
        rpr.insert(0, rfonts)
    for attr in ("ascii", "hAnsi", "eastAsia", "cs"):
        rfonts.set(qn(f"w:{attr}"), "Microsoft YaHei")


def shade(cell, fill: str) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    shd = tcpr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tcpr.append(shd)
    shd.set(qn("w:fill"), fill)


def cell_margin(cell, *, top=110, start=145, bottom=110, end=145) -> None:
    tcpr = cell._tc.get_or_add_tcPr()
    mar = tcpr.first_child_found_in("w:tcMar")
    if mar is None:
        mar = OxmlElement("w:tcMar")
        tcpr.append(mar)
    for side, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        el = mar.find(qn(f"w:{side}"))
        if el is None:
            el = OxmlElement(f"w:{side}")
            mar.append(el)
        el.set(qn("w:w"), str(value))
        el.set(qn("w:type"), "dxa")


def run_text(paragraph, value: str, *, bold=False, code=False) -> None:
    run = paragraph.add_run(value)
    run.bold = bold
    if code:
        run.font.name = "Consolas"
        run.font.size = Pt(9.2)
        run.font.color.rgb = BLUE


TOKEN = re.compile(r"(\*\*[^*]+\*\*|`[^`]+`|\[[^]]+\]\([^)]+\))")


def add_inline(paragraph, value: str) -> None:
    offset = 0
    for match in TOKEN.finditer(value):
        if match.start() > offset:
            run_text(paragraph, value[offset:match.start()])
        token = match.group()
        if token.startswith("**"):
            run_text(paragraph, token[2:-2], bold=True)
        elif token.startswith("`"):
            run_text(paragraph, token[1:-1], code=True)
        else:
            label, url = token[1:].split("](", 1)
            run_text(paragraph, f"{label}（{url[:-1]}）")
        offset = match.end()
    if offset < len(value):
        run_text(paragraph, value[offset:])


def add_code(document, lines: list[str]) -> None:
    for line in lines:
        p = document.add_paragraph(style="Code")
        p.paragraph_format.keep_together = True
        run_text(p, line, code=True)


def add_table(document, rows: list[list[str]]) -> None:
    table = document.add_table(rows=0, cols=2)
    table.autofit = False
    table.columns[0].width = Inches(2.12)
    table.columns[1].width = Inches(4.70)
    table.style = "Table Grid"
    for index, row in enumerate(rows):
        cells = table.add_row().cells
        cells[0].width = Inches(2.12)
        cells[1].width = Inches(4.70)
        for cell, value in zip(cells, row):
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            cell_margin(cell)
            p = cell.paragraphs[0]
            p.paragraph_format.space_after = Pt(0)
            p.paragraph_format.line_spacing = 1.17
            add_inline(p, value)
            if index == 0:
                shade(cell, PALE)
                for run in p.runs:
                    run.bold = True
            elif index % 2 == 0:
                shade(cell, "F8FAFB")
        if index == 0:
            trpr = table.rows[index]._tr.get_or_add_trPr()
            tbl_header = OxmlElement("w:tblHeader")
            tbl_header.set(qn("w:val"), "true")
            trpr.append(tbl_header)
    document.add_paragraph().paragraph_format.space_after = Pt(0)


def main() -> None:
    source = SOURCE.read_text(encoding="utf-8-sig").splitlines()
    doc = Document()
    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(0.72)
    section.bottom_margin = Inches(0.68)
    section.left_margin = Inches(0.83)
    section.right_margin = Inches(0.83)

    styles = doc.styles
    normal = styles["Normal"]
    set_font(normal, size=10.5)
    normal.paragraph_format.space_after = Pt(7)
    normal.paragraph_format.line_spacing = 1.30

    title = styles["Title"]
    set_font(title, size=19.5, bold=True)
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.keep_with_next = True

    heading = styles["Heading 1"]
    set_font(heading, size=13.5, bold=True)
    heading.paragraph_format.space_before = Pt(17)
    heading.paragraph_format.space_after = Pt(8)
    heading.paragraph_format.keep_with_next = True

    code_style = styles.add_style("Code", 1)
    set_font(code_style, size=9.2, color=BLUE)
    code_style.paragraph_format.left_indent = Inches(0.18)
    code_style.paragraph_format.space_after = Pt(2)
    code_style.paragraph_format.line_spacing = 1.12

    bullet = styles["List Number"]
    set_font(bullet, size=10.5)
    bullet.paragraph_format.space_after = Pt(5)
    bullet.paragraph_format.line_spacing = 1.25

    index = 0
    while index < len(source):
        line = source[index].strip()
        if not line:
            index += 1
            continue
        if line.startswith("# "):
            doc.add_paragraph(line[2:], style="Title")
        elif line.startswith("## "):
            doc.add_paragraph(line[3:], style="Heading 1")
        elif line.startswith("```powershell"):
            lines = []
            index += 1
            while index < len(source) and not source[index].startswith("```"):
                lines.append(source[index])
                index += 1
            add_code(doc, lines)
        elif line.startswith("| "):
            rows = []
            while index < len(source) and source[index].startswith("| "):
                parts = [part.strip() for part in source[index].strip().strip("|").split("|")]
                if not all(re.fullmatch(r":?-+:?", part) for part in parts):
                    rows.append(parts)
                index += 1
            add_table(doc, rows)
            continue
        elif re.match(r"^\d+\. ", line):
            p = doc.add_paragraph(style="List Number")
            add_inline(p, re.sub(r"^\d+\. ", "", line))
        else:
            p = doc.add_paragraph()
            add_inline(p, line)
        index += 1

    doc.core_properties.title = "安护智链 CareGuard 新同事 Windows 启动指南"
    doc.core_properties.subject = "网页与本地工程启动说明"
    doc.core_properties.keywords = "CareGuard, 安护智链, Windows, 启动指南"
    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    main()
